"""Unit tests for `app.ingestion.extraction` (Issue #3, Slice 3.4). Pure
(extension, bytes) -> ExtractedDocument / ExtractionError — no database,
no HTTP, no real filesystem. See `tests/test_document_processing.py` for
the end-to-end HTTP-level coverage of the `/process` endpoint that calls
into this module.
"""

from __future__ import annotations

import io
import zipfile
from unittest.mock import patch

import docx
import pytest
from pypdf import PdfWriter

from app.ingestion import extraction


def _blank_pdf(*, pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_with_headings() -> bytes:
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("body text one")
    document.add_heading("Second", level=2)
    document.add_paragraph("body text two")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_with_single_member(name: str, content: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, content)
    return buffer.getvalue()


# --- PDF ----------------------------------------------------------------


def test_pdf_valid_blank_pages_extracts_with_correct_page_count() -> None:
    result = extraction.extract(extension=".pdf", content=_blank_pdf(pages=2))
    assert result.page_count == 2
    assert len(result.sections) == 2
    assert result.sections[0].page == 1
    assert result.sections[1].page == 2


def test_pdf_with_no_extractable_text_does_not_crash() -> None:
    # A structurally valid PDF with blank pages has no text content stream
    # -- extraction must return empty text per page, never raise.
    result = extraction.extract(extension=".pdf", content=_blank_pdf(pages=1))
    assert result.sections[0].text == ""


def test_pdf_malformed_bytes_raise_extraction_error() -> None:
    with pytest.raises(extraction.ExtractionError):
        extraction.extract(extension=".pdf", content=b"%PDF-1.4\nnot really a pdf")


def test_pdf_completely_non_pdf_bytes_raise_extraction_error() -> None:
    with pytest.raises(extraction.ExtractionError):
        extraction.extract(extension=".pdf", content=b"just some random bytes, not a pdf at all")


def test_pdf_page_count_limit_is_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "_MAX_PDF_PAGES", 2)
    with pytest.raises(extraction.ExtractionError, match="too many pages"):
        extraction.extract(extension=".pdf", content=_blank_pdf(pages=3))


def test_pdf_single_page_parser_exception_is_normalized() -> None:
    # Simulates pypdf itself raising while reading one specific page's
    # content stream -- a real malformed-content-stream PDF is difficult
    # to construct deterministically, so the library seam is patched
    # directly to prove this failure path is caught and normalized
    # rather than propagating a raw pypdf exception.
    with (
        patch("pypdf._page.PageObject.extract_text", side_effect=RuntimeError("boom")),
        pytest.raises(extraction.ExtractionError, match="page 1"),
    ):
        extraction.extract(extension=".pdf", content=_blank_pdf(pages=1))


# --- DOCX -----------------------------------------------------------------


def test_docx_valid_splits_sections_on_headings() -> None:
    result = extraction.extract(extension=".docx", content=_docx_with_headings())
    assert result.page_count is None
    headings = [section.heading for section in result.sections]
    assert headings == ["Title", "Second"]
    assert "body text one" in result.sections[0].text
    assert "body text two" in result.sections[1].text


def test_docx_malformed_zip_raises_extraction_error() -> None:
    with pytest.raises(extraction.ExtractionError, match="corrupt"):
        extraction.extract(extension=".docx", content=b"not a zip file at all")


@pytest.mark.parametrize(
    "member_name",
    [
        "../evil.txt",
        "..\\evil.txt",
        "/etc/passwd",
        "a/../../evil.txt",
    ],
)
def test_docx_archive_traversal_member_name_rejected(member_name: str) -> None:
    content = _docx_with_single_member(member_name, b"x")
    with pytest.raises(extraction.ExtractionError, match="unsafe internal file path"):
        extraction.extract(extension=".docx", content=content)


def test_docx_member_count_limit_is_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "_MAX_DOCX_MEMBER_COUNT", 3)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for i in range(5):
            archive.writestr(f"part{i}.xml", b"x")
    with pytest.raises(extraction.ExtractionError, match="too many internal parts"):
        extraction.extract(extension=".docx", content=buffer.getvalue())


def test_docx_per_member_size_limit_is_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "_MAX_DOCX_MEMBER_UNCOMPRESSED_BYTES", 1000)
    content = _docx_with_single_member("big.xml", b"a" * 2000)
    with pytest.raises(extraction.ExtractionError, match="oversized internal part"):
        extraction.extract(extension=".docx", content=content)


def test_docx_total_uncompressed_size_limit_is_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "_MAX_DOCX_MEMBER_UNCOMPRESSED_BYTES", 10_000)
    monkeypatch.setattr(extraction, "_MAX_DOCX_TOTAL_UNCOMPRESSED_BYTES", 1500)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("a.xml", b"a" * 900)
        archive.writestr("b.xml", b"b" * 900)
    with pytest.raises(extraction.ExtractionError, match="unsafe size"):
        extraction.extract(extension=".docx", content=buffer.getvalue())


def test_docx_real_zip_bomb_style_member_rejected_at_default_thresholds() -> None:
    # A highly-compressible member whose *declared* uncompressed size
    # exceeds the real (unmodified) per-member cap, while its actual
    # compressed footprint is tiny -- exactly the shape of a zip-bomb
    # attack. Proves the check is against declared uncompressed size, not
    # the file's on-disk/compressed size.
    content = _docx_with_single_member("bomb.xml", b"\x00" * 60_000_000)
    assert len(content) < 100_000  # compressed footprint stayed tiny
    with pytest.raises(extraction.ExtractionError, match="oversized internal part"):
        extraction.extract(extension=".docx", content=content)


def test_docx_parser_failure_after_safety_check_is_normalized() -> None:
    # Passes the ZIP-level safety check (a plausible member name/size) but
    # is not a real OOXML document -- python-docx itself must fail, and
    # that failure must be normalized, not propagated raw.
    content = _docx_with_single_member("word/document.xml", b"not real xml content")
    with pytest.raises(extraction.ExtractionError, match="corrupt"):
        extraction.extract(extension=".docx", content=content)


# --- TXT / Markdown / CSV --------------------------------------------------


def test_txt_valid_utf8_decodes() -> None:
    result = extraction.extract(extension=".txt", content="hello, world — café".encode())
    assert result.sections[0].text == "hello, world — café"
    assert result.page_count is None


def test_txt_invalid_byte_sequence_does_not_crash() -> None:
    result = extraction.extract(extension=".txt", content=b"valid text \xff\xfe invalid bytes")
    assert "�" in result.sections[0].text


def test_markdown_splits_on_top_level_headings() -> None:
    content = b"# Heading One\n\nfirst body\n\n## Heading Two\n\nsecond body\n"
    result = extraction.extract(extension=".md", content=content)
    assert [s.heading for s in result.sections] == ["Heading One", "Heading Two"]
    assert "first body" in result.sections[0].text
    assert "second body" in result.sections[1].text


def test_markdown_with_no_headings_is_a_single_section() -> None:
    result = extraction.extract(extension=".md", content=b"just plain content, no headings")
    assert len(result.sections) == 1
    assert result.sections[0].heading is None


def test_markdown_empty_document_does_not_crash() -> None:
    result = extraction.extract(extension=".md", content=b"")
    assert len(result.sections) == 1
    assert result.sections[0].text == ""


def test_csv_normal_content_is_rendered() -> None:
    result = extraction.extract(extension=".csv", content=b"col1,col2\n1,2\n3,4\n")
    assert "col1, col2" in result.sections[0].text
    assert "1, 2" in result.sections[0].text


def test_csv_field_exceeding_size_limit_raises_extraction_error() -> None:
    content = b"a" * 200_000
    with pytest.raises(extraction.ExtractionError, match="malformed"):
        extraction.extract(extension=".csv", content=content)


# --- dispatch / output budget ----------------------------------------------


def test_unsupported_extension_raises_extraction_error() -> None:
    with pytest.raises(extraction.ExtractionError):
        extraction.extract(extension=".exe", content=b"anything")


def test_extracted_text_budget_is_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "_MAX_EXTRACTED_TEXT_BYTES", 10)
    content = b"this text is much longer than ten bytes"
    result = extraction.extract(extension=".txt", content=content)
    assert len(result.sections[0].text.encode("utf-8")) <= 10


def test_full_text_joins_sections() -> None:
    document = extraction.ExtractedDocument(
        sections=[
            extraction.ExtractedSection(text="first"),
            extraction.ExtractedSection(text="second"),
        ]
    )
    assert document.full_text == "first\n\nsecond"
