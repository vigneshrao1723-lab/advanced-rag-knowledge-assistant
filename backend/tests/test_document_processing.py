"""HTTP-level tests for the document processing endpoint (Issue #3, Slice
3.4):

    POST /api/v1/workspaces/{workspace_id}/documents/{document_id}/process

Real Postgres, real Redis, real filesystem (an isolated per-test root via
a `get_storage_provider` dependency override) — no mocks, matching this
project's established convention (see `tests/test_document_upload.py`).
Every document here is uploaded through the real upload endpoint first,
then processed through this endpoint -- exercising the two slices
together, the way a real client would.
"""

from __future__ import annotations

import asyncio
import io
import time
import uuid
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

import docx
import httpx
import pytest
import redis
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession

from app.core.audit import AuditEvent
from app.core.redis_client import build_redis_client, get_redis_client
from app.core.redis_keys import rate_limit_key
from app.ingestion import extraction
from app.models.audit_log import AuditLog
from app.models.document import Document, DocumentStatus
from app.services.storage_provider import (
    LocalStorage,
    StorageError,
    StorageProvider,
    get_storage_provider,
)
from tests.conftest import csrf_headers

_PASSWORD = "correct horse battery staple"


def _blank_pdf(*, pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_with_heading() -> bytes:
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("hello world body text")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_PDF_BYTES = _blank_pdf(pages=2)
_MALFORMED_PDF_BYTES = b"%PDF-1.4\n" + b"garbage" * 5  # passes the upload signature check only
_DOCX_BYTES = _docx_with_heading()
_MALFORMED_DOCX_BYTES = b"PK\x03\x04" + b"\x00" * 40  # passes the upload signature check only
_TXT_BYTES = b"hello world, this is a plain text document."
_MD_BYTES = b"# Heading\n\nSome **markdown** text."
_CSV_BYTES = b"col1,col2\n1,2\n3,4\n"
_CSV_FIELD_TOO_LARGE_BYTES = b"a" * 200_000  # exceeds csv's default field_size_limit


def _unique_email() -> str:
    return f"user-{uuid.uuid4().hex[:12]}@example.com"


def _register(client: TestClient, email: str | None = None) -> dict[str, Any]:
    email = email or _unique_email()
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": _PASSWORD},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text
    result: dict[str, Any] = response.json()
    return result


def _login(client: TestClient, email: str) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": _PASSWORD},
        headers=csrf_headers(client),
    )
    assert response.status_code == 200, response.text


@pytest.fixture(autouse=True)
def _isolated_storage(app: FastAPI, tmp_path: Path) -> Iterator[None]:
    app.dependency_overrides[get_storage_provider] = lambda: LocalStorage(root=str(tmp_path))
    yield
    del app.dependency_overrides[get_storage_provider]


@pytest.fixture
def storage_root(tmp_path: Path) -> Path:
    return tmp_path


def _create_workspace(client: TestClient, name: str = "Acme") -> dict[str, Any]:
    response = client.post(
        "/api/v1/workspaces", json={"name": name}, headers=csrf_headers(client)
    )
    assert response.status_code == 201, response.text
    result: dict[str, Any] = response.json()
    return result


def _add_member(client: TestClient, workspace_id: str, email: str, role: str) -> None:
    response = client.post(
        f"/api/v1/workspaces/{workspace_id}/members",
        json={"email": email, "role": role},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text


def _upload(
    client: TestClient,
    workspace_id: str,
    *,
    filename: str,
    content: bytes,
    content_type: str,
) -> dict[str, Any]:
    response = client.post(
        f"/api/v1/workspaces/{workspace_id}/documents",
        files={"file": (filename, content, content_type)},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201, response.text
    result: dict[str, Any] = response.json()
    return result


def _process(client: TestClient, workspace_id: str, document_id: str) -> Any:
    return client.post(
        f"/api/v1/workspaces/{workspace_id}/documents/{document_id}/process",
        headers=csrf_headers(client),
    )


def _audit_rows(db_session: DbSession, *, event_type: str, workspace_id: str) -> list[AuditLog]:
    return list(
        db_session.execute(
            select(AuditLog).where(
                AuditLog.event_type == event_type,
                AuditLog.workspace_id == uuid.UUID(workspace_id),
            )
        ).scalars()
    )


# --- per-format success: transitions to PARSED -----------------------------


def test_pdf_processes_to_parsed_with_correct_page_count(
    client: TestClient, db_session: DbSession
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "PARSED"
    assert body["page_count"] == 2
    assert body["failure_reason"] is None
    assert "storage_key" not in body

    row = db_session.get(Document, uuid.UUID(document["id"]))
    assert row is not None
    assert row.status == DocumentStatus.PARSED
    assert row.processing_started_at is not None
    assert row.processing_completed_at is not None


def test_docx_processes_to_parsed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.docx",
        content=_DOCX_BYTES,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "PARSED"
    assert body["page_count"] is None


def test_txt_processes_to_parsed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client, workspace["id"], filename="a.txt", content=_TXT_BYTES, content_type="text/plain"
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "PARSED"


def test_markdown_processes_to_parsed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client, workspace["id"], filename="a.md", content=_MD_BYTES, content_type="text/markdown"
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "PARSED"


def test_csv_processes_to_parsed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client, workspace["id"], filename="a.csv", content=_CSV_BYTES, content_type="text/csv"
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "PARSED"


def test_successful_processing_emits_exactly_one_document_parsed_audit_event(
    client: TestClient, db_session: DbSession
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200

    rows = _audit_rows(
        db_session, event_type=AuditEvent.DOCUMENT_PARSED, workspace_id=workspace["id"]
    )
    assert len(rows) == 1
    row = rows[0]
    assert row.event_metadata is not None
    assert row.event_metadata["document_id"] == document["id"]
    assert row.event_metadata["page_count"] == 2


# --- malformed content: transitions to FAILED, never crashes ---------------


def test_malformed_pdf_transitions_to_failed_not_a_500(
    client: TestClient, db_session: DbSession
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="bad.pdf",
        content=_MALFORMED_PDF_BYTES,
        content_type="application/pdf",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert body["failure_reason"]
    assert "storage_key" not in body
    assert str(document["id"]) not in body["failure_reason"]

    row = db_session.get(Document, uuid.UUID(document["id"]))
    assert row is not None
    assert row.status == DocumentStatus.FAILED
    assert row.failure_reason is not None


def test_malformed_docx_transitions_to_failed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="bad.docx",
        content=_MALFORMED_DOCX_BYTES,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "FAILED"


def test_docx_archive_traversal_attempt_transitions_to_failed(client: TestClient) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("../evil.txt", b"x")
    content = buffer.getvalue()

    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="evil.docx",
        content=content,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert "evil" not in body["failure_reason"]  # no internal archive member name leaked


def test_malformed_csv_transitions_to_failed(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="huge.csv",
        content=_CSV_FIELD_TOO_LARGE_BYTES,
        content_type="text/csv",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "FAILED"


def test_failed_processing_emits_document_parsing_failed_audit_event(
    client: TestClient, db_session: DbSession
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="bad.pdf",
        content=_MALFORMED_PDF_BYTES,
        content_type="application/pdf",
    )

    _process(client, workspace["id"], document["id"])

    rows = _audit_rows(
        db_session, event_type=AuditEvent.DOCUMENT_PARSING_FAILED, workspace_id=workspace["id"]
    )
    assert len(rows) == 1
    assert rows[0].event_metadata is not None
    assert rows[0].event_metadata["document_id"] == document["id"]


# --- storage read failure ---------------------------------------------------


class _ReadFailsStorage:
    """Wraps a real LocalStorage but forces read() to fail with
    StorageError -- exactly the exception type StorageProvider
    implementations guarantee on a genuine failure, never a raw OSError."""

    def __init__(self, inner: StorageProvider) -> None:
        self._inner = inner

    def save(self, *, key: str, content: bytes) -> None:
        self._inner.save(key=key, content=content)

    def read(self, *, key: str) -> bytes:
        raise StorageError(f"simulated read failure for key {key!r}")

    def delete(self, *, key: str) -> None:
        self._inner.delete(key=key)

    def exists(self, *, key: str) -> bool:
        return self._inner.exists(key=key)


def test_storage_read_failure_transitions_to_failed_without_leaking_the_key(
    app: FastAPI, client: TestClient, storage_root: Path
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    app.dependency_overrides[get_storage_provider] = lambda: _ReadFailsStorage(
        LocalStorage(root=str(storage_root))
    )
    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert document["id"] not in body["failure_reason"]
    assert "simulated read failure" not in body["failure_reason"]


# --- unexpected parser exception --------------------------------------------


def test_unexpected_extraction_exception_transitions_to_failed_not_a_500(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    def _boom(*, extension: str, content: bytes) -> extraction.ExtractedDocument:
        raise RuntimeError("totally unexpected failure")

    monkeypatch.setattr(extraction, "extract", _boom)

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "FAILED"
    assert "totally unexpected failure" not in body["failure_reason"]


# --- crash-safety: PROCESSING is committed before extraction runs ----------


def test_processing_transition_is_committed_before_extraction_is_attempted(
    client: TestClient, db_session: DbSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Extraction runs via asyncio.to_thread() (see document_service.py's
    # _read_and_extract() docstring -- a synchronous, potentially
    # CPU-bound call inside an async endpoint would otherwise block this
    # project's single uvicorn process for every other concurrent
    # request), so the extraction spy below executes on a worker thread,
    # not the thread that owns `db_session`. SQLAlchemy Sessions are not
    # thread-safe, so this test deliberately does NOT query the database
    # from inside the spy (unlike a same-thread version of this test
    # could) -- instead it records the *relative order* of "a commit
    # happened" vs. "extraction started" via plain list.append() calls,
    # which is safe across threads under the GIL, and proves the same
    # property: PROCESSING must be committed strictly before extraction
    # is even attempted.
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    order: list[str] = []
    real_extract = extraction.extract
    real_commit = db_session.commit

    def _tracking_commit() -> None:
        order.append("commit")
        real_commit()

    def _spy_extract(*, extension: str, content: bytes) -> extraction.ExtractedDocument:
        order.append("extract_started")
        return real_extract(extension=extension, content=content)

    monkeypatch.setattr(db_session, "commit", _tracking_commit)
    monkeypatch.setattr(extraction, "extract", _spy_extract)

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200, response.text

    # order[0] is the PROCESSING-transition commit; "extract_started"
    # (and every later commit, for mark_parsed/mark_failed) must come
    # strictly after it -- if the commit happened only after extraction,
    # a crash during extraction could leave the document looking like it
    # was never touched (still UPLOADED) rather than honestly PROCESSING.
    assert order[0] == "commit"
    assert order.index("extract_started") > order.index("commit")


# --- resource exhaustion: a slow extraction must not block other requests --


@pytest.mark.asyncio
async def test_slow_extraction_does_not_block_unrelated_concurrent_requests(
    app: FastAPI, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # This project runs one uvicorn process with no --workers (see
    # infra/docker/backend.Dockerfile's entrypoint), so if extraction ran
    # synchronously inside the async endpoint, a single slow/pathological
    # document would stall every other concurrent request this process is
    # serving -- not just the caller's own -- for the full duration of
    # parsing. document_service._read_and_extract() runs via
    # asyncio.to_thread() specifically to prevent this. Proven here with
    # real concurrency (httpx.AsyncClient over an in-process ASGI
    # transport, sharing the same event loop the app itself runs on) and
    # an absolute shared clock, not per-request-relative timings, which
    # would hide exactly this kind of event-loop-blocking bug.
    app.dependency_overrides[get_storage_provider] = lambda: LocalStorage(root=str(tmp_path))

    def _slow_extract(*, extension: str, content: bytes) -> extraction.ExtractedDocument:
        time.sleep(1.0)  # simulates real CPU-bound parser work
        return extraction.ExtractedDocument(sections=[], page_count=1)

    monkeypatch.setattr(extraction, "extract", _slow_extract)

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as async_client:
        csrf_response = await async_client.get("/api/v1/auth/csrf")
        csrf_token = csrf_response.cookies.get("csrf_token")
        headers = {"X-CSRF-Token": csrf_token or ""}

        register_response = await async_client.post(
            "/api/v1/auth/register",
            json={"email": _unique_email(), "password": _PASSWORD},
            headers=headers,
        )
        assert register_response.status_code == 201, register_response.text

        workspace_response = await async_client.post(
            "/api/v1/workspaces", json={"name": "concurrency-check"}, headers=headers
        )
        assert workspace_response.status_code == 201, workspace_response.text
        workspace_id = workspace_response.json()["id"]

        upload_response = await async_client.post(
            f"/api/v1/workspaces/{workspace_id}/documents",
            files={"file": ("a.pdf", _PDF_BYTES, "application/pdf")},
            headers=headers,
        )
        assert upload_response.status_code == 201, upload_response.text
        document_id = upload_response.json()["id"]

        t_zero = time.monotonic()

        async def slow_process() -> None:
            response = await async_client.post(
                f"/api/v1/workspaces/{workspace_id}/documents/{document_id}/process",
                headers=headers,
            )
            assert response.status_code == 200, response.text

        async def fast_health() -> float:
            await asyncio.sleep(0.2)
            response = await async_client.get("/api/v1/health")
            assert response.status_code == 200
            return time.monotonic() - t_zero

        _, fast_finish = await asyncio.gather(slow_process(), fast_health())

    # The health check was scheduled to fire at t=0.2s. If extraction were
    # blocking the event loop, it wouldn't complete until the ~1s slow
    # extraction finished too. A generous threshold (well under the 1s
    # extraction time) avoids test flakiness from ordinary scheduling
    # jitter while still failing hard if the event loop was blocked.
    assert fast_finish < 0.6


# --- authorization / lifecycle -----------------------------------------------


def test_unauthenticated_process_rejected(client_factory: Callable[[], TestClient]) -> None:
    anonymous = client_factory()
    response = _process(anonymous, str(uuid.uuid4()), str(uuid.uuid4()))
    assert response.status_code == 401


def test_viewer_process_rejected(client_factory: Callable[[], TestClient]) -> None:
    owner = client_factory()
    _register(owner)
    workspace = _create_workspace(owner)
    document = _upload(
        owner, workspace["id"], filename="a.pdf", content=_PDF_BYTES, content_type="application/pdf"
    )

    viewer_email = _unique_email()
    _register(client_factory(), viewer_email)
    _add_member(owner, workspace["id"], viewer_email, "VIEWER")

    viewer = client_factory()
    _login(viewer, viewer_email)

    response = _process(viewer, workspace["id"], document["id"])
    assert response.status_code == 403


def test_non_member_process_rejected_with_404_not_403(
    client_factory: Callable[[], TestClient],
) -> None:
    owner = client_factory()
    _register(owner)
    workspace = _create_workspace(owner)
    document = _upload(
        owner, workspace["id"], filename="a.pdf", content=_PDF_BYTES, content_type="application/pdf"
    )

    outsider = client_factory()
    _register(outsider)

    response = _process(outsider, workspace["id"], document["id"])
    assert response.status_code == 404


def test_cross_workspace_document_id_rejected_with_404(
    client_factory: Callable[[], TestClient],
) -> None:
    owner_a = client_factory()
    _register(owner_a)
    workspace_a = _create_workspace(owner_a, name="A")
    document = _upload(
        owner_a,
        workspace_a["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    owner_b = client_factory()
    _register(owner_b)
    workspace_b = _create_workspace(owner_b, name="B")

    # owner_b is a real member of workspace_b, but the document belongs to
    # workspace_a -- must not be reachable through workspace_b's ID.
    response = _process(owner_b, workspace_b["id"], document["id"])
    assert response.status_code == 404


def test_nonexistent_document_id_rejected_with_404(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    response = _process(client, workspace["id"], str(uuid.uuid4()))
    assert response.status_code == 404


def test_processing_an_already_parsed_document_is_rejected_with_409(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    first = _process(client, workspace["id"], document["id"])
    assert first.status_code == 200
    assert first.json()["status"] == "PARSED"

    second = _process(client, workspace["id"], document["id"])
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "document_already_processed"


def test_a_failed_document_can_be_retried(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    document = _upload(
        client,
        workspace["id"],
        filename="bad.pdf",
        content=_MALFORMED_PDF_BYTES,
        content_type="application/pdf",
    )

    first = _process(client, workspace["id"], document["id"])
    assert first.status_code == 200
    assert first.json()["status"] == "FAILED"

    # Retriable: a FAILED document is not permanently stuck. The
    # underlying bytes are still malformed, so it fails again -- but as a
    # 200/FAILED outcome, never a 409, proving the retry path itself
    # works even though this particular retry doesn't succeed.
    second = _process(client, workspace["id"], document["id"])
    assert second.status_code == 200
    assert second.json()["status"] == "FAILED"


# --- rate limiting ------------------------------------------------------


@pytest.fixture
def unreachable_redis_client() -> Iterator[redis.Redis]:
    client = build_redis_client(
        "redis://localhost:1/0", socket_timeout=0.05, socket_connect_timeout=0.05
    )
    yield client
    client.close()


def test_process_rate_limit_dimensions_are_recorded_in_redis(
    client: TestClient, redis_client: redis.Redis
) -> None:
    _register(client)
    workspace = _create_workspace(client)
    me = client.get("/api/v1/users/me")
    assert me.status_code == 200
    user_id = me.json()["id"]
    document = _upload(
        client,
        workspace["id"],
        filename="a.pdf",
        content=_PDF_BYTES,
        content_type="application/pdf",
    )

    response = _process(client, workspace["id"], document["id"])
    assert response.status_code == 200

    ip_key = rate_limit_key("document_process", "ip", "testclient")
    user_key = rate_limit_key("document_process", "user", user_id)
    assert redis_client.exists(ip_key) == 1
    assert redis_client.exists(user_key) == 1


def test_process_rate_limit_enforced_at_threshold(client: TestClient) -> None:
    _register(client)
    workspace = _create_workspace(client)
    # A single FAILED document, reprocessed repeatedly: FAILED is
    # retriable (`_REPROCESSABLE_STATUSES`), so each call genuinely
    # re-attempts extraction and returns 200/FAILED again, letting this
    # test exhaust the *process* rate limit without also tripping the
    # separate upload rate limit (which is also 20/60s) by uploading 25
    # distinct documents.
    document = _upload(
        client,
        workspace["id"],
        filename="bad.pdf",
        content=_MALFORMED_PDF_BYTES,
        content_type="application/pdf",
    )

    allowed = 0
    last_status = None
    for _ in range(25):
        response = _process(client, workspace["id"], document["id"])
        last_status = response.status_code
        if response.status_code == 200:
            allowed += 1
        else:
            break

    assert allowed == 20
    assert last_status == 429


def test_process_redis_unavailable_falls_back_and_still_enforces_the_limit(
    app: FastAPI,
    client_factory: Callable[[], TestClient],
    unreachable_redis_client: redis.Redis,
) -> None:
    app.dependency_overrides[get_redis_client] = lambda: unreachable_redis_client
    test_client = client_factory()
    _register(test_client)
    workspace = _create_workspace(test_client)
    document = _upload(
        test_client,
        workspace["id"],
        filename="bad.pdf",
        content=_MALFORMED_PDF_BYTES,
        content_type="application/pdf",
    )

    allowed = 0
    last_status = None
    for _ in range(25):
        response = _process(test_client, workspace["id"], document["id"])
        last_status = response.status_code
        if response.status_code == 200:
            allowed += 1
        else:
            break

    assert allowed == 20
    assert last_status == 429


@pytest.fixture
def redis_client() -> Iterator[redis.Redis]:
    client = get_redis_client()
    assert client is not None
    yield client
